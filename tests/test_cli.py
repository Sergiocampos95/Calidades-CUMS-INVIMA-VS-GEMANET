"""Pruebas de la CLI (`gemma-cum-loader`).

Existen porque la CLI estuvo rota sin que nadie lo notara: quedo importando
`procesar_malla`, una funcion que desaparecio cuando el pipeline dejo de
recibir una "malla" externa y paso a armar los candidatos desde el catalogo
INVIMA. Sin una prueba que la ejecute, un entry point declarado en
pyproject.toml puede romperse en un refactor y solo enterarse el usuario
final. El test mas importante de este archivo es el mas tonto: que importe y
que corra de punta a punta.
"""

from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from gemma_cum_loader.cli import main

# Mismo esquema minimo que usa test_invima_reader (29 columnas del catalogo
# oficial); aca solo hace falta que el lector lo acepte y produzca 1 candidato.
COLUMNAS_INVIMA = [
    "EXPEDIENTE", "PRODUCTO", "TITULAR", "REGISTRO SANITARIO",
    "FECHA EXPEDICION", "FECHA VENCIMIENTO", "ESTADO REGISTRO",
    "EXPEDIENTE CUM", "CONSECUTIVO", "CANTIDAD CUM", "DESCRIPCION COMERCIAL",
    "ESTADO CUM", "FECHA ACTIVO", "FECHA INACTIVO", "MUESTRA MEDICA",
    "UNIDAD", "ATC", "DESCRIPCION_ATC", "VIA ADMINISTRACION",
    "CONCENTRACION", "PRINCIPIO ACTIVO", "UNIDAD MEDIDA", "CANTIDAD",
    "UNIDAD REFERENCIA", "FORMA FARMACEUTICA", "NOMBRE ROL", "TIPO ROL",
    "MODALIDAD", "IUM",
]

FILA_INVIMA = [
    10815, "PROD A", "TIT A", "INVIMA 123", None, None, "Vigente", 10815, 3, 1,
    "DESC A", "Activo", None, None, "No", "MG", "N06AB03", "ANTIDEPRESIVO",
    "ORAL", "20 MG", "FLUOXETINA", "MG", 20, "MG", "CAPSULA",
    "LAB X", "FABRICANTE", "NACIONAL", None,
]


@pytest.fixture
def invima_xlsx(tmp_path):
    ruta = tmp_path / "ListadoCodigoUnicoVigentes.xlsx"
    wb = openpyxl.Workbook()
    hoja = wb.active
    hoja.title = "Vigentes"
    # El export real de INVIMA trae 6 filas de preambulo antes del encabezado
    # -- el lector las salta (header=6), asi que la prueba las reproduce.
    for _ in range(6):
        hoja.append([None] * len(COLUMNAS_INVIMA))
    hoja.append(COLUMNAS_INVIMA)
    hoja.append(FILA_INVIMA)
    wb.save(ruta)
    return ruta


@pytest.fixture
def gemanet_xlsx(tmp_path):
    """Reporte de Gemma Net con un codigo que NO esta en el INVIMA de prueba
    -- asi la auditoria produce al menos una fila y candidatos_creacion deja
    al menos un candidato nuevo."""
    ruta = tmp_path / "reporte_gemanet.xlsx"
    pd.DataFrame(
        [{"CODIGO_INTERNO": "99999-1", "DESCRIPCION": "LEGADO", "ACTIVO": "Si"}]
    ).to_excel(ruta, index=False)
    return ruta


def test_cli_se_puede_importar_y_expone_main():
    """Guardia contra la regresion exacta que rompio la CLI: un import roto
    a nivel de modulo (pyproject declara `gemma_cum_loader.cli:main`)."""
    assert callable(main)


def test_candidatos_escribe_salida_y_retorna_cero(invima_xlsx, gemanet_xlsx, tmp_path):
    salida = tmp_path / "candidatos.xlsx"
    codigo = main([
        "candidatos",
        "--invima", str(invima_xlsx),
        "--gemma-net", str(gemanet_xlsx),
        "--salida", str(salida),
    ])
    assert codigo == 0
    assert salida.exists()


def test_auditoria_escribe_salida_con_estado_coherencia(invima_xlsx, gemanet_xlsx, tmp_path):
    salida = tmp_path / "auditoria.xlsx"
    codigo = main([
        "auditoria",
        "--invima", str(invima_xlsx),
        "--gemma-net", str(gemanet_xlsx),
        "--salida", str(salida),
    ])
    assert codigo == 0
    # guardar_reporte abre una hoja por valor de ESTADO_COHERENCIA -- el
    # codigo 99999-1 no esta en el INVIMA de prueba, asi que cae ahi.
    hojas = openpyxl.load_workbook(salida).sheetnames
    assert "sin_correspondencia_invima" in hojas


def test_auditoria_sigue_sin_los_3_auxiliares_opcionales(invima_xlsx, gemanet_xlsx, tmp_path):
    """Vencidos / Renovacion / Otros Estados son opcionales e independientes:
    su ausencia degrada la clasificacion, nunca tumba la corrida."""
    salida = tmp_path / "auditoria_sin_aux.xlsx"
    assert main([
        "auditoria",
        "--invima", str(invima_xlsx),
        "--gemma-net", str(gemanet_xlsx),
        "--salida", str(salida),
    ]) == 0
    assert salida.exists()


def test_archivo_inexistente_falla_con_mensaje_claro_no_traceback(gemanet_xlsx, tmp_path):
    """argparse.error -> SystemExit(2), no un traceback crudo de pandas."""
    with pytest.raises(SystemExit) as exc_info:
        main([
            "auditoria",
            "--invima", str(tmp_path / "no_existe.xlsx"),
            "--gemma-net", str(gemanet_xlsx),
            "--salida", str(tmp_path / "x.xlsx"),
        ])
    assert exc_info.value.code == 2


def test_exige_elegir_origen_del_catalogo_invima(gemanet_xlsx, tmp_path):
    """--invima y --api son mutuamente excluyentes y uno es obligatorio."""
    with pytest.raises(SystemExit):
        main(["auditoria", "--gemma-net", str(gemanet_xlsx), "--salida", str(tmp_path / "x.xlsx")])


def test_sin_subcomando_falla(tmp_path):
    with pytest.raises(SystemExit):
        main([])
