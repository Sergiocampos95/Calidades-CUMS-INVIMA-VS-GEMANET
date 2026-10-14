from pathlib import Path

import openpyxl
import pytest

from gemma_cum_loader.ingesta.invima_reader import (
    _es_formato_portal,
    catalogo_activo_fabricante,
    leer_catalogo_invima,
    leer_catalogo_invima_otros_estados,
    leer_catalogo_invima_renovacion,
    leer_catalogo_invima_vencidos,
    mapa_ium,
    universo_codigos_internos,
)

COLUMNAS = [
    "EXPEDIENTE", "PRODUCTO", "TITULAR", "REGISTRO SANITARIO",
    "FECHA EXPEDICION", "FECHA VENCIMIENTO", "ESTADO REGISTRO",
    "EXPEDIENTE CUM", "CONSECUTIVO", "CANTIDAD CUM", "DESCRIPCION COMERCIAL",
    "ESTADO CUM", "FECHA ACTIVO", "FECHA INACTIVO", "MUESTRA MEDICA",
    "UNIDAD", "ATC", "DESCRIPCION_ATC", "VIA ADMINISTRACION",
    "CONCENTRACION", "PRINCIPIO ACTIVO", "UNIDAD MEDIDA", "CANTIDAD",
    "UNIDAD REFERENCIA", "FORMA FARMACEUTICA", "NOMBRE ROL", "TIPO ROL",
    "MODALIDAD", "IUM",
]

FILAS = [
    # mismo expediente-consecutivo, dos roles -> se deduplica al filtrar FABRICANTE
    [10815, "PROD A", "TIT A", "INVIMA 123", None, None, "Vigente", 10815, 3, 1,
     "SIN DATO EN ORIGEN", "Activo", None, None, "No", "MG", "N06AB03", "ANTIDEPRESIVO",
     "ORAL", "SIN DATO", "FLUOXETINA CLORHIDRATO", "MG", 20, "MG", "CAPSULA DURA",
     "LAB X", "FABRICANTE", "NACIONAL", None],
    [10815, "PROD A", "TIT A", "INVIMA 123", None, None, "Vigente", 10815, 3, 1,
     "SIN DATO EN ORIGEN", "Activo", None, None, "No", "MG", "N06AB03", "ANTIDEPRESIVO",
     "ORAL", "SIN DATO", "FLUOXETINA CLORHIDRATO", "MG", 20, "MG", "CAPSULA DURA",
     "LAB X", "IMPORTADOR", "NACIONAL", None],
    # presentacion vigente pero con CUM inactivo (caso de las 19.542 filas del diagnostico)
    [20147, "PROD B", "TIT B", "INVIMA 456", None, None, "Vigente", 20147, 1, 1,
     "DESC B", "Inactivo", None, None, "No", "G", "A10BA02", "ANTIDIABETICO",
     "ORAL", "DESC B CONC", "METFORMINA", "G", 500, "MG", "TABLETA",
     "LAB Y", "FABRICANTE", "NACIONAL", None],
    # fila con IUM (cobertura ~9% del catalogo)
    [30111, "PROD C", "TIT C", "INVIMA 789", None, None, "Vigente", 30111, 1, 1,
     "DESC C", "Activo", None, None, "No", "MG", "N06AB03", "ANTIDEPRESIVO",
     "ORAL", "DESC C CONC", "SERTRALINA", "MG", 50, "MG", "TABLETA",
     "LAB Z", "FABRICANTE", "NACIONAL", "2B1040311000103"],
    # medicamento combinado: mismo fabricante, dos filas por dos principios activos
    # (caso real verificado: 57% de los "duplicados" del archivo son esto, no
    # fabricantes en competencia -- no se deben colapsar a una sola fila)
    [40222, "PROD D", "TIT D", "INVIMA 999", None, None, "Vigente", 40222, 1, 1,
     "DESC D", "Activo", None, None, "No", "MG", "R05CA10", "ANTITUSIVO COMBINADO",
     "ORAL", "DESC D CONC", "CLORFENIRAMINA MALEATO", "MG", 4, "MG", "JARABE",
     "LAB W", "FABRICANTE", "NACIONAL", None],
    [40222, "PROD D", "TIT D", "INVIMA 999", None, None, "Vigente", 40222, 1, 1,
     "DESC D", "Activo", None, None, "No", "MG", "R05CA10", "ANTITUSIVO COMBINADO",
     "ORAL", "DESC D CONC", "DEXTROMETORFANO BROMHIDRATO", "MG", 15, "MG", "JARABE",
     "LAB W", "FABRICANTE", "NACIONAL", None],
]


def _crear_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vigente"
    for _ in range(6):  # filas 1-6: titulo y banda "CUM", el encabezado real va en fila 7
        ws.append([None] * len(COLUMNAS))
    ws.append(COLUMNAS)
    for fila in FILAS:
        ws.append(fila)
    ruta = tmp_path / "ListadoCodigoUnicoVigentes2022.xlsx"
    wb.save(ruta)
    return ruta


def test_leer_catalogo_construye_codigo_interno(tmp_path):
    df = leer_catalogo_invima(_crear_xlsx(tmp_path))
    assert "CODIGO_INTERNO" in df.columns
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_activo_fabricante_deduplica_por_rol_y_excluye_inactivos(tmp_path):
    df = leer_catalogo_invima(_crear_xlsx(tmp_path))
    activos = catalogo_activo_fabricante(df)
    # 40222-1 aparece dos veces: no se colapsa, son dos principios activos
    # distintos del mismo medicamento combinado, no un fabricante duplicado
    assert list(activos["CODIGO_INTERNO"]) == ["10815-3", "30111-1", "40222-1", "40222-1"]


def test_universo_codigos_internos_no_pierde_ni_colapsa_de_mas(tmp_path):
    df = leer_catalogo_invima(_crear_xlsx(tmp_path))
    universo = universo_codigos_internos(catalogo_activo_fabricante(df))
    # como set, el medicamento combinado (dos filas) cuenta una sola vez para
    # pertenencia, sin necesidad de tiebreak ni de perder ningun principio activo
    assert universo == {"10815-3", "30111-1", "40222-1"}


def test_vigente_pero_cum_inactivo_no_esta_en_universo_activo_pero_si_en_crudo(tmp_path):
    df = leer_catalogo_invima(_crear_xlsx(tmp_path))
    activos = catalogo_activo_fabricante(df)
    assert "20147-1" not in set(activos["CODIGO_INTERNO"])
    assert "20147-1" in set(df["CODIGO_INTERNO"])


def test_mapa_ium_solo_incluye_filas_con_ium(tmp_path):
    df = leer_catalogo_invima(_crear_xlsx(tmp_path))
    ium = mapa_ium(df)
    assert list(ium["CODIGO_INTERNO"]) == ["30111-1"]


# ---- leer_catalogo_invima_vencidos / _renovacion / _otros_estados: sin
# archivo real descargado para confirmar el nombre de hoja de ninguno de
# los 3, se prueba la misma cascada de tolerancia que ya funciona para la
# malla de referencia (_resolver_hoja_malla) ----

def _crear_xlsx_auxiliar(
    tmp_path: Path, nombre_archivo: str, nombre_hoja: str, hojas_extra: list[str] | None = None
) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja
    for _ in range(6):
        ws.append([None] * len(COLUMNAS))
    ws.append(COLUMNAS)
    for fila in FILAS:
        ws.append(fila)
    for extra in hojas_extra or []:
        wb.create_sheet(extra)
    ruta = tmp_path / nombre_archivo
    wb.save(ruta)
    return ruta


# ---- leer_catalogo_invima (Vigentes): tolerancia de nombre de hoja --
# bug real de produccion (2026-08-19, "Worksheet named 'Vigente' not
# found"): esta funcion tenia el nombre de hoja hardcoded sin la misma
# tolerancia que ya tenian los otros 3 lectores, y fallo con un archivo
# Vigentes real cuya hoja no se llamaba exactamente "Vigente" ----

def test_leer_vigentes_con_hoja_candidata_plural(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoVigentes.xlsx", "Vigentes")
    df = leer_catalogo_invima(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vigentes_con_hoja_unica_sin_nombre_reconocible(tmp_path):
    # el caso real que motivo este fix: la hoja no coincide con ningun
    # candidato exacto ni con el fragmento "vigen", pero es la unica hoja
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoVigentes.xlsx", "Hoja1")
    df = leer_catalogo_invima(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vigentes_falla_con_mensaje_claro_si_hay_varias_hojas_ambiguas(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoVigentes.xlsx", "Hoja1", hojas_extra=["Hoja2"])
    with pytest.raises(ValueError) as exc_info:
        leer_catalogo_invima(ruta)
    assert "Hoja1" in str(exc_info.value)
    assert "Hoja2" in str(exc_info.value)


def _crear_xlsx_vencidos(tmp_path: Path, nombre_hoja: str, hojas_extra: list[str] | None = None) -> Path:
    return _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoVencidos.xlsx", nombre_hoja, hojas_extra)


def test_leer_vencidos_con_hoja_candidata_singular(tmp_path):
    df = leer_catalogo_invima_vencidos(_crear_xlsx_vencidos(tmp_path, "Vencido"))
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vencidos_con_hoja_candidata_plural(tmp_path):
    df = leer_catalogo_invima_vencidos(_crear_xlsx_vencidos(tmp_path, "Vencidos"))
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vencidos_con_formato_actual_sin_filas_preliminares_ni_ium(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    encabezados = [col.replace(" ", "").replace("_", "") for col in COLUMNAS[:-1]]
    encabezados[8] = "CONSECUTIVOCUM"
    ws.append(encabezados)
    for fila in FILAS:
        ws.append(fila[:-1])
    ruta = tmp_path / "Codigo_Unico_Medicamentos_Vencidos.xlsx"
    wb.save(ruta)

    df = leer_catalogo_invima_vencidos(ruta)

    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}
    assert "ESTADO_REGISTRO" in df.columns
    assert "IUM" not in df.columns


def test_leer_vencidos_con_hoja_renombrada_por_coincidencia_de_substring(tmp_path):
    # nombre real del archivo puede variar entre copias -- ninguno de los
    # candidatos exactos calza, pero "venc" si aparece y es la unica hoja
    # de datos, igual que _resolver_hoja_malla tolera variantes de "plantilla"
    df = leer_catalogo_invima_vencidos(_crear_xlsx_vencidos(tmp_path, "Vencidos INVIMA"))
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vencidos_con_hoja_unica_sin_nombre_reconocible(tmp_path):
    df = leer_catalogo_invima_vencidos(_crear_xlsx_vencidos(tmp_path, "Hoja1"))
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_vencidos_falla_con_mensaje_claro_si_hay_varias_hojas_ambiguas(tmp_path):
    ruta = _crear_xlsx_vencidos(tmp_path, "Hoja1", hojas_extra=["Hoja2"])
    with pytest.raises(ValueError) as exc_info:
        leer_catalogo_invima_vencidos(ruta)
    assert "Hoja1" in str(exc_info.value)
    assert "Hoja2" in str(exc_info.value)


def test_leer_vencidos_no_filtra_por_estado_cum_ni_tipo_rol(tmp_path):
    # a diferencia de catalogo_activo_fabricante (logica de negocio de
    # Vigentes), Vencidos solo necesita el set de CODIGO_INTERNO tal cual --
    # la fila con ESTADO_CUM=Inactivo (20147-1) debe seguir presente
    df = leer_catalogo_invima_vencidos(_crear_xlsx_vencidos(tmp_path, "Vencido"))
    assert "20147-1" in set(df["CODIGO_INTERNO"])
    # y el medicamento combinado conserva sus dos filas (dos principios activos)
    assert list(df["CODIGO_INTERNO"]).count("40222-1") == 2


# ---- leer_catalogo_invima_renovacion: mismo patron tolerante, candidatos
# de hoja distintos ("Tramite de Renovacion"/"Renovacion"/"Renovaciones") ----

def test_leer_renovacion_con_hoja_candidata_exacta(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Tramite de Renovacion")
    df = leer_catalogo_invima_renovacion(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_renovacion_con_hoja_candidata_alterna(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Renovaciones")
    df = leer_catalogo_invima_renovacion(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_renovacion_con_hoja_en_tramite_nombre_real_del_archivo_de_invima(tmp_path):
    # nombre real confirmado contra el archivo de Renovacion del usuario
    # (2026-08-19, ListadoCodigoUnicoRenovacion2022.xlsx) -- no coincide con
    # ningun candidato original ni con el fragmento "renov"
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "En tramite")
    df = leer_catalogo_invima_renovacion(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_renovacion_con_hoja_renombrada_por_coincidencia_de_substring(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Renovacion INVIMA")
    df = leer_catalogo_invima_renovacion(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_renovacion_con_hoja_unica_sin_nombre_reconocible(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Hoja1")
    df = leer_catalogo_invima_renovacion(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_renovacion_falla_con_mensaje_claro_si_hay_varias_hojas_ambiguas(tmp_path):
    ruta = _crear_xlsx_auxiliar(
        tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Hoja1", hojas_extra=["Hoja2"]
    )
    with pytest.raises(ValueError) as exc_info:
        leer_catalogo_invima_renovacion(ruta)
    assert "Hoja1" in str(exc_info.value)
    assert "Hoja2" in str(exc_info.value)


def test_leer_renovacion_no_filtra_por_estado_cum_ni_tipo_rol(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoRenovacion.xlsx", "Tramite de Renovacion")
    df = leer_catalogo_invima_renovacion(ruta)
    assert "20147-1" in set(df["CODIGO_INTERNO"])
    assert list(df["CODIGO_INTERNO"]).count("40222-1") == 2


# ---- leer_catalogo_invima_otros_estados: mismo patron tolerante, ademas
# ESTADO_REGISTRO debe sobrevivir intacto (la auditoria lo necesita, el
# dataset es heterogeneo -- Cancelado/Suspendido/Inactivo/etc.) ----

def test_leer_otros_estados_con_hoja_candidata_exacta(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Otros Estados")
    df = leer_catalogo_invima_otros_estados(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_otros_estados_con_hoja_renombrada_por_coincidencia_de_substring(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Otro Estado INVIMA")
    df = leer_catalogo_invima_otros_estados(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_otros_estados_con_hoja_unica_sin_nombre_reconocible(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Hoja1")
    df = leer_catalogo_invima_otros_estados(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}


def test_leer_otros_estados_falla_con_mensaje_claro_si_hay_varias_hojas_ambiguas(tmp_path):
    ruta = _crear_xlsx_auxiliar(
        tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Hoja1", hojas_extra=["Hoja2"]
    )
    with pytest.raises(ValueError) as exc_info:
        leer_catalogo_invima_otros_estados(ruta)
    assert "Hoja1" in str(exc_info.value)
    assert "Hoja2" in str(exc_info.value)


def test_leer_otros_estados_no_filtra_por_estado_cum_ni_tipo_rol(tmp_path):
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Otros Estados")
    df = leer_catalogo_invima_otros_estados(ruta)
    assert "20147-1" in set(df["CODIGO_INTERNO"])
    assert list(df["CODIGO_INTERNO"]).count("40222-1") == 2


def test_leer_otros_estados_conserva_estado_registro_real_por_fila(tmp_path):
    # el dataset es heterogeneo (Cancelado/Suspendido/Inactivo/etc.) -- la
    # auditoria necesita el valor real, no solo saber que "esta en otros
    # estados". Todas las filas del fixture comparten ESTADO REGISTRO=
    # "Vigente" (ver FILAS), asi que basta confirmar que la columna
    # sobrevive con ese valor real en vez de perderse/vaciarse.
    ruta = _crear_xlsx_auxiliar(tmp_path, "ListadoCodigoUnicoOtrosEstados.xlsx", "Otros Estados")
    df = leer_catalogo_invima_otros_estados(ruta)
    assert "ESTADO_REGISTRO" in df.columns
    assert set(df["ESTADO_REGISTRO"]) == {"Vigente"}


# ---- formato nuevo del portal de datos.gov.co (visto 2026-09-28): hoja
# "Data", encabezado en la fila 1, columnas estilo API (registrosanitario,
# tiporol...) en vez del Excel clasico (fila 7, "REGISTRO SANITARIO",
# "TIPO ROL"). Se detecta por el nombre de columnas, nunca por el nombre de
# hoja ni por cuantas columnas trae -- ver docstring de _es_formato_portal.

COLUMNAS_API = [
    "expediente", "producto", "titular", "registrosanitario",
    "fechaexpedicion", "fechavencimiento", "estadoregistro",
    "expedientecum", "consecutivocum", "cantidadcum", "descripcioncomercial",
    "estadocum", "fechaactivo", "fechainactivo", "muestramedica",
    "unidad", "atc", "descripcionatc", "viaadministracion",
    "concentracion", "principioactivo", "unidadmedida", "cantidad",
    "unidadreferencia", "formafarmaceutica", "nombrerol", "tiporol",
    "modalidad", "ium",
]


def _crear_xlsx_formato_portal(tmp_path: Path, nombre_archivo: str, columnas: list[str], filas: list) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"  # el portal siempre nombra la hoja "Data", sin decir el dataset
    ws.append(columnas)  # encabezado en la fila 1, no en la 7
    for fila in filas:
        ws.append(fila)
    ruta = tmp_path / nombre_archivo
    wb.save(ruta)
    return ruta


def test_es_formato_portal_reconoce_columnas_estilo_api():
    assert _es_formato_portal(COLUMNAS_API)


def test_es_formato_portal_no_confunde_una_banda_de_titulo_clasica():
    # las primeras filas del Excel clasico son titulo/banda "CUM", no encabezado real
    assert not _es_formato_portal([None] * len(COLUMNAS))


def test_leer_catalogo_reconoce_formato_nuevo_del_portal_completo(tmp_path):
    ruta = _crear_xlsx_formato_portal(tmp_path, "CUM_vigente_20260928.xlsx", COLUMNAS_API, FILAS)
    df = leer_catalogo_invima(ruta)
    assert set(df["CODIGO_INTERNO"]) == {"10815-3", "20147-1", "30111-1", "40222-1"}
    assert "TIPO_ROL" in df.columns
    assert "ESTADO_CUM" in df.columns


def test_leer_vencidos_reconoce_formato_nuevo_del_portal(tmp_path):
    ruta = _crear_xlsx_formato_portal(tmp_path, "Codigo_Unico_Medicamentos_Vencidos.xlsx", COLUMNAS_API, FILAS)
    df = leer_catalogo_invima_vencidos(ruta)
    assert "40222-1" in set(df["CODIGO_INTERNO"])


def test_leer_catalogo_formato_portal_incompleto_no_arma_codigo_interno_si_falta_consecutivo(tmp_path):
    # caso real 2026-09-28: Vigentes vino en este formato con solo 10 de las
    # 29 columnas, sin "consecutivocum" -- no se inventa un CONSECUTIVO, y
    # sin el, CODIGO_INTERNO tampoco se arma (degradacion explicita: el
    # dato no esta, no se adivina).
    columnas_incompletas = [
        "expediente", "principioactivo", "cantidad", "unidadmedida",
        "unidadreferencia", "formafarmaceutica", "registrosanitario",
        "fechaexpedicion", "fechavencimiento", "estadoregistro",
    ]
    fila = [20038383, "MONTELUKAST SODICO", 10, "mg", "TABLETA", "TABLETA RECUBIERTA",
            "INVIMA 2021M-0013144-R1", "04/30/2012", None, "Vigente"]
    ruta = _crear_xlsx_formato_portal(tmp_path, "CUM_vigente_incompleto.xlsx", columnas_incompletas, [fila])
    df = leer_catalogo_invima(ruta)
    assert "EXPEDIENTE" in df.columns
    assert "CODIGO_INTERNO" not in df.columns
