import datetime as dt

import openpyxl
import pytest

from gemma_cum_loader.armado.reglas_negocio import (
    CLASIFICADO_VALORES_VALIDOS,
    NIVELES_SERVICIO_VALIDOS,
    derivar_reglas_negocio,
    leer_malla_referencia,
)

ENCABEZADOS = [
    "EXPEDIENTE",
    "CLASIFICADO(1:SI, 2:NO, 3.MEDICAMENTO ANCESTRAL, 4. PLAN MEDICINAL)",
    "CRES(SI/NO)",
    "GENERA COPAGO RS(SI/NO)",
    "GENERA COPAGO RC(SI/NO)",
    "GENERA CUOTA MODERADORA(SI/NO)",
    "AUTOMATICO(SI/NO)",
    "CAMBIO CANTIDAD(SI/NO)",
    "VALOR",
    "REGULADO(SI/NO)",
    "VALOR REGULADO",
    "RESTRINGIDO(SI/NO)",
    "EDAD MÍN. AÑOS",
    "EDAD MÁX. AÑOS",
    "MÁX. VECES DÍA",
    "MÁX. VECES MES",
    "MÁX. VECES AÑO",
    "MÁX. VECES VIDA",
    "TIEMPO LÍMITE DÍAS",
    "GRUPO MEDICAMENTO",
    "CÓDIGO NIVEL SERVICIO",
    "POSOLOGÍA",
    "DÍAS",
    "POS(SI/NO)",
    "CÓDIGO INTERNO MODELO SERVICIO",
]


def _fila_constante(expediente=500, **overrides):
    base = {
        "EXPEDIENTE": expediente,
        "CLASIFICADO(1:SI, 2:NO, 3.MEDICAMENTO ANCESTRAL, 4. PLAN MEDICINAL)": "NO",
        "CRES(SI/NO)": "NO",
        "GENERA COPAGO RS(SI/NO)": "NO",
        "GENERA COPAGO RC(SI/NO)": "NO",
        "GENERA CUOTA MODERADORA(SI/NO)": "SI",
        "AUTOMATICO(SI/NO)": "NO",
        "CAMBIO CANTIDAD(SI/NO)": "SI",
        "VALOR": 0,
        "REGULADO(SI/NO)": "NO",
        "VALOR REGULADO": 0,
        "RESTRINGIDO(SI/NO)": "NO",
        "EDAD MÍN. AÑOS": 0,
        "EDAD MÁX. AÑOS": 100,
        "MÁX. VECES DÍA": 0,
        "MÁX. VECES MES": 0,
        "MÁX. VECES AÑO": 0,
        "MÁX. VECES VIDA": 0,
        "TIEMPO LÍMITE DÍAS": 0,
        "GRUPO MEDICAMENTO": 1,
        "CÓDIGO NIVEL SERVICIO": "Nivel 1",
        "POSOLOGÍA": 0,
        "DÍAS": 0,
        "POS(SI/NO)": "SI",
        "CÓDIGO INTERNO MODELO SERVICIO": 39,
    }
    base.update(overrides)
    return [base[c] for c in ENCABEZADOS]


def _crear_malla_xlsx(tmp_path, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "plantilla (2)"
    ws.append(ENCABEZADOS)
    for fila in filas:
        ws.append(fila)
    ruta = tmp_path / "malla_referencia.xlsx"
    wb.save(ruta)
    return ruta


def _crear_catalogo_modelo_servicio(tmp_path, codigos_validos):
    ruta = tmp_path / "modelo_servicio.csv"
    contenido = "codigo,texto\n" + "\n".join(f"{c},MODELO {c}" for c in codigos_validos)
    ruta.write_text(contenido + "\n", encoding="utf-8")
    return ruta


def test_campos_100_por_ciento_constantes_no_generan_advertencia(tmp_path):
    filas = [_fila_constante(expediente=e) for e in range(10)]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39, 40])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert reglas.valor("CLASIFICADO") == "NO"
    assert reglas.valor("GENERA_CUOTA_MODERADORA") == "SI"
    assert reglas.valor("EDAD_MAXIMA") == 100
    assert reglas.valor("GRUPO_MEDICAMENTO") == 1
    assert reglas.valor("CODIGO_NIVEL_SERVICIO") == "Nivel 1"
    assert "CLASIFICADO" not in reglas.advertencias
    assert "GENERA_CUOTA_MODERADORA" not in reglas.advertencias


def test_campo_constante_con_una_sola_excepcion_advierte(tmp_path):
    # a diferencia de POS/modelo de servicio, estos campos SI son ruido si
    # aparece una unica excepcion -- pero igual hay que advertir, no ocultarla
    filas = [_fila_constante(expediente=e) for e in range(9)] + [
        _fila_constante(expediente=999, **{"CRES(SI/NO)": "SI"})
    ]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert reglas.valor("CRES") == "NO"
    assert "CRES" in reglas.advertencias
    assert "90.0%" in reglas.advertencias["CRES"]


def test_activo_y_fecha_inicio_son_logicos_no_dependen_de_la_malla(tmp_path):
    filas = [_fila_constante(expediente=e) for e in range(3)]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert reglas.valor("ACTIVO") == "SI"
    assert reglas.valor("FECHA_INICIO") == dt.date.today().strftime("%Y/%m/%d")


def test_fecha_fin_queda_vacia_por_ser_el_dato_correcto(tmp_path):
    filas = [_fila_constante(expediente=e) for e in range(3)]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert reglas.valor("FECHA_FIN") == ""


def test_pos_se_clasifica_por_expediente_no_por_mayoria_global(tmp_path):
    # expediente 500 siempre POS=SI en sus presentaciones, expediente 600
    # siempre POS=NO -- son dos medicamentos distintos, cada uno con su
    # propia clasificacion oficial, no "una mayoria global" que se les aplique
    # a ambos por igual
    filas = [
        _fila_constante(expediente=500, **{"POS(SI/NO)": "SI"}),
        _fila_constante(expediente=500, **{"POS(SI/NO)": "SI"}),
        _fila_constante(expediente=600, **{"POS(SI/NO)": "NO"}),
    ]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    clasif_500 = reglas.clasificar("POS", 500)
    assert clasif_500.cierta is True
    assert clasif_500.valor == "SI"

    clasif_600 = reglas.clasificar("POS", 600)
    assert clasif_600.cierta is True
    assert clasif_600.valor == "NO"


def test_expediente_con_pos_mixto_no_es_cierto(tmp_path):
    # caso real raro (2 de 9.799 expedientes en la malla real): el mismo
    # expediente trae POS distinto entre sus propias presentaciones -- no se
    # puede confiar en un solo valor
    filas = [
        _fila_constante(expediente=700, **{"POS(SI/NO)": "SI"}),
        _fila_constante(expediente=700, **{"POS(SI/NO)": "NO"}),
    ]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    clasif = reglas.clasificar("POS", 700)
    assert clasif.cierta is False
    assert "valores distintos" in clasif.motivo


def test_expediente_sin_precedente_no_es_cierto(tmp_path):
    # un candidato de un expediente que nunca aparecio en la malla de
    # referencia -- no hay ninguna fuente que diga su POS, se marca pendiente
    filas = [_fila_constante(expediente=500)]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    clasif = reglas.clasificar("POS", 999999)
    assert clasif.cierta is False
    assert "no tiene ninguna presentacion previa" in clasif.motivo


def test_codigo_modelo_servicio_invalido_contra_catalogo_real_genera_advertencia(tmp_path):
    filas = [
        _fila_constante(expediente=e, **{"CÓDIGO INTERNO MODELO SERVICIO": 999}) for e in range(5)
    ]
    malla = leer_malla_referencia(_crear_malla_xlsx(tmp_path, filas))
    # el catalogo real no tiene el codigo 999
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39, 40])

    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert "CODIGO_INTERNO_MODELO_SERVICIO" in reglas.advertencias
    assert "5 expediente(s)" in reglas.advertencias["CODIGO_INTERNO_MODELO_SERVICIO"]


def test_columna_faltante_en_la_malla_genera_advertencia_no_rompe(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "plantilla (2)"
    ws.append(["OTRA COLUMNA"])
    ws.append(["x"])
    ruta = tmp_path / "malla_incompleta.xlsx"
    wb.save(ruta)
    catalogo = _crear_catalogo_modelo_servicio(tmp_path, [39])

    malla = leer_malla_referencia(ruta)
    reglas = derivar_reglas_negocio(malla, ruta_catalogo_modelo_servicio=catalogo)

    assert "CLASIFICADO" in reglas.advertencias
    assert reglas.valor("CLASIFICADO") == ""
    assert "POS" in reglas.advertencias
    assert reglas.clasificar("POS", 500).cierta is False


def test_valores_validos_documentados_coinciden_con_tablas_de_referencia_reales():
    assert CLASIFICADO_VALORES_VALIDOS == ["SI", "NO", "Medicamento Ancestral", "Planta Medicinal"]
    assert NIVELES_SERVICIO_VALIDOS == ["Nivel 1", "Nivel 2", "Nivel 3", "Nivel 4"]


def test_lee_malla_con_hoja_renombrada_por_coincidencia(tmp_path):
    # caso real: ValueError "Worksheet named 'plantilla (2)' not found" -- una
    # copia distinta del archivo trae la misma hoja con otro nombre. Si solo
    # hay una hoja que contiene "plantilla" en el nombre, se usa esa en vez
    # de fallar.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    ws.append(ENCABEZADOS)
    ruta = tmp_path / "malla_renombrada.xlsx"
    wb.save(ruta)

    malla = leer_malla_referencia(ruta)
    assert list(malla.columns) == ENCABEZADOS


def test_lee_malla_con_hoja_unica_sin_importar_el_nombre(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(ENCABEZADOS)
    ruta = tmp_path / "malla_hoja_unica.xlsx"
    wb.save(ruta)

    malla = leer_malla_referencia(ruta)
    assert list(malla.columns) == ENCABEZADOS


def test_falla_con_mensaje_claro_si_hay_varias_hojas_ambiguas(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws.append(ENCABEZADOS)
    wb.create_sheet("Hoja2")
    ruta = tmp_path / "malla_ambigua.xlsx"
    wb.save(ruta)

    with pytest.raises(ValueError) as exc_info:
        leer_malla_referencia(ruta)
    assert "Hoja1" in str(exc_info.value) and "Hoja2" in str(exc_info.value)
